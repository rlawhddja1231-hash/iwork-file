#!/usr/bin/env python3
"""iwork.py — Apple iWork file tool (Pages, Numbers, Keynote) for Claude.

Lets Claude read Apple .pages / .numbers / .key files without a Mac, edit
Numbers spreadsheets, find-and-replace text in Pages and Keynote, build new
Numbers files, pull out pictures, and convert all three to Office formats
(Excel with working formulas, Word, PowerPoint) or Markdown.

Usage
  python3 iwork.py info    FILE                       # what is inside (type, sheets/slides, size, lock)
  python3 iwork.py read    FILE [--out text.md]       # full content as Markdown (tables, slides, notes)
  python3 iwork.py preview FILE [--out folder]        # pull the built-in preview pictures (slide thumbnails)
  python3 iwork.py convert FILE OUT [--values-only]   # .numbers -> .xlsx (live formulas)/.csv/.md
                                                      # .pages   -> .docx/.md/.txt
                                                      # .key     -> .pptx/.docx/.md/.txt
                                                      # old iWork '09 files with a stored PDF -> .pdf
  python3 iwork.py set     FILE.numbers A1=value ... [--sheet S] [--table T] [--out new.numbers]
  python3 iwork.py replace FILE FIND REPLACE [FIND2 REPLACE2 ...] [--dry] [--out new.FILE]
                                                      # Numbers cells, or Pages/Keynote text (2013+ files)
  python3 iwork.py images  FILE [--out folder] [--all] # save the pictures/media used in the document
  python3 iwork.py new     DATA.xlsx|DATA.csv [--out new.numbers]
  Add --password PW to open a password-protected Numbers file.

Needs: pip install numbers-parser keynote-parser openpyxl python-docx python-pptx
Built on numbers-parser (MIT, masaccio) and keynote-parser (MIT, psobot).
Not affiliated with or endorsed by Apple Inc.
"""
import argparse
import csv
import datetime as dt
import gzip
import io
import json
import os
import re
import sys
import warnings
import zipfile
from pathlib import Path

warnings.filterwarnings("ignore")
warnings.showwarning = lambda *args, **kwargs: None   # library chatter is not useful to the user

KINDS = {".pages": "pages", ".numbers": "numbers", ".key": "keynote"}


def die(msg, code=1):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(code)


def need(mod, pipname):
    try:
        return __import__(mod)
    except ImportError:
        die(f"missing package — run: pip install --break-system-packages {pipname}")


# ───────────────────────── package access (zip file or Mac folder bundle) ─────────────────────────
class Package:
    """An iWork file is either a zip file or (when copied from a Mac) a folder."""

    def __init__(self, path):
        self.path = Path(path)
        if not self.path.exists():
            die(f"file not found: {path}")
        self.kind = KINDS.get(self.path.suffix.lower())
        if self.path.is_dir():
            self.names = sorted(str(p.relative_to(self.path)).replace(os.sep, "/")
                                for p in self.path.rglob("*") if p.is_file())
            self._zip = None
        else:
            try:
                self._zip = zipfile.ZipFile(self.path)
            except zipfile.BadZipFile:
                die("this file is damaged or is not an iWork document (not a valid zip package)")
            self.names = self._zip.namelist()
        # Some exports put everything inside one top folder, e.g. "Doc.pages/Index/..."
        self.prefix = ""
        if not any(n.startswith(("Index/", "index.", "Index.zip", "presentation.apxl")) for n in self.names):
            tops = {n.split("/")[0] for n in self.names if "/" in n}
            for t in tops:
                if any(n.startswith((t + "/Index/", t + "/index.", t + "/Index.zip", t + "/presentation.apxl"))
                       for n in self.names):
                    self.prefix = t + "/"
                    break
        # Older Numbers/Pages (2013-2015) keep IWA files inside an inner Index.zip
        self._inner = None
        if self.has("Index.zip") and not any(n.startswith(self.prefix + "Index/") for n in self.names):
            try:
                self._inner = zipfile.ZipFile(io.BytesIO(self.read("Index.zip")))
            except zipfile.BadZipFile:
                die("this file is damaged (its inner content package is broken)")

    def has(self, name):
        return (self.prefix + name) in self.names

    def read(self, name):
        full = self.prefix + name
        if self._zip is not None:
            try:
                return self._zip.read(full)
            except (NotImplementedError, RuntimeError):
                die("this file is password-protected or uses packing this tool cannot open. "
                    "Ask the owner to remove the password (File > Set Password) or export it to Word/Excel/PowerPoint.")
        return (self.path / full).read_bytes()

    def list(self):
        out = [n[len(self.prefix):] for n in self.names if n.startswith(self.prefix) and not n.endswith("/")]
        if self._inner is not None:
            out += self._inner.namelist()
        return out

    def read_any(self, name):
        if self._inner is not None and name in self._inner.namelist():
            return self._inner.read(name)
        return self.read(name)

    @property
    def format(self):
        names = self.list()
        if any(n.endswith(".iwph") for n in names) or ".iwpv2" in names:
            return "locked"
        if any(n.startswith("Index/") and n.endswith(".iwa") for n in names):
            return "iwa"
        if any(n in ("index.xml", "index.xml.gz", "index.apxl", "index.apxl.gz", "presentation.apxl",
                     "presentation.apxl.gz") for n in names):
            return "xml"
        return "unknown"


# ───────────────────────── IWA (2013+) decoding through keynote-parser ─────────────────────────
class IWA:
    """All archives of a modern iWork file, by identifier."""

    def __init__(self, pkg):
        need("keynote_parser", "keynote-parser")
        from keynote_parser.codec import IWAFile
        self.pkg = pkg
        self.objs = {}      # id -> dict
        self.where = {}     # id -> iwa file name
        self.order = []     # ids in file order
        self.failed = []    # parts that could not be decoded
        self.datarefs = {}  # iwa file -> media ids its objects use
        for n in pkg.list():
            if not n.endswith(".iwa"):
                continue
            try:
                d = IWAFile.from_buffer(pkg.read_any(n), n).to_dict()
            except Exception:  # noqa: BLE001
                self.failed.append(n)
                continue
            for ch in d.get("chunks", []):
                for a in ch.get("archives", []):
                    ident = str(a.get("header", {}).get("identifier"))
                    for mi in a.get("header", {}).get("messageInfos", []):
                        for dref in mi.get("dataReferences", []) or []:
                            self.datarefs.setdefault(n, set()).add(str(dref))
                    objs = a.get("objects") or [{}]
                    self.objs[ident] = objs[0]
                    self.where[ident] = n
                    self.order.append(ident)

    def get(self, ref):
        if isinstance(ref, dict):
            ref = ref.get("identifier")
        return self.objs.get(str(ref), {})

    def of_type(self, t):
        return [(i, self.objs[i]) for i in self.order if self.objs[i].get("_pbtype") == t]

    def style_info(self, ref):
        """Walk a style's parent chain and collect its name, built-in id and list label types."""
        info = {"name": "", "ident": "", "labels": None}
        seen = 0
        o = self.get(ref) if ref else {}
        while o and seen < 10:
            sup = o.get("super", {})
            info["name"] = info["name"] or sup.get("name") or ""
            info["ident"] = info["ident"] or sup.get("styleIdentifier") or ""
            if info["labels"] is None and o.get("labelTypes"):
                info["labels"] = o["labelTypes"]
            o = self.get(sup.get("parent")) if sup.get("parent") else {}
            seen += 1
        return info

    def style_name(self, ref):
        return self.style_info(ref)["name"]

    def data_file(self, ident):
        """Map a media data id to its file inside Data/."""
        ident = str(ident)
        if not hasattr(self, "_datas"):
            self._datas = {}
            for _, o in self.of_type("TSP.PackageMetadata"):
                for x in o.get("datas", []):
                    self._datas[str(x.get("identifier"))] = x
        x = self._datas.get(ident, {})
        names = self.pkg.list()
        for cand in (x.get("fileName"), x.get("preferredFileName")):
            if cand and ("Data/" + cand) in names:
                return "Data/" + cand
        for n in names:
            if n.startswith("Data/") and re.search(rf"-{ident}\.\w+$", n):
                return n
        return None


CLEAN = {"\u2028": "\n", "\u2029": "\n", "\x04": "\n\n---\n\n", "\x05": "", "\x0c": "\n\n---\n\n",
         "\ufffc": "", "\x0e": "", "\x0b": "\n"}


def clean_text(s):
    for a, b in CLEAN.items():
        s = s.replace(a, b)
    return "".join(c for c in s if c in "\n\t" or c.isprintable())


def storage_paragraphs(iwa, st):
    """Split a TSWP.StorageArchive into (style name, list marker, text) paragraphs."""
    text = "".join(st.get("text", []))
    if not text:
        return []
    # characterIndex counts UTF-16 units; map to Python indices
    u16 = []
    for i, c in enumerate(text):
        u16.append(i)
        if ord(c) > 0xFFFF:
            u16.append(i)
    u16.append(len(text))

    def py(ix):
        return u16[min(ix, len(u16) - 1)]

    def table(name):
        return [(py(e.get("characterIndex", 0)), e) for e in st.get(name, {}).get("entries", [])]

    para_styles = table("tableParaStyle")
    list_styles = table("tableListStyle")
    para_data = table("tableParaData")

    def at(tab, i):
        cur = None
        for ix, e in tab:
            if ix <= i:
                cur = e
            else:
                break
        return cur

    out, start = [], 0
    for m in re.finditer(r"[\n\u2029]", text + "\n"):
        seg = text[start:m.start()]
        ps = at(para_styles, start)
        ls = at(list_styles, start)
        pd = at(para_data, start)
        pinfo = iwa.style_info(ps.get("object")) if ps else {"name": "", "ident": ""}
        linfo = iwa.style_info(ls.get("object")) if ls else {"labels": None}
        level = (pd or {}).get("first", 0) or 0
        labels = linfo.get("labels") or []
        label = labels[min(level, len(labels) - 1)] if labels else "kNone"
        lkind = {"kString": "bullet", "kImage": "bullet", "kNumber": "number"}.get(label, "")
        # the built-in id ends with the English style name (e.g. "text-0-paragraphstyle-Title")
        ident_tail = re.sub(r"^.*paragraphstyle-?", "", pinfo.get("ident") or "") if "paragraphstyle-" in (pinfo.get("ident") or "") else ""
        out.append((ident_tail or pinfo.get("name", ""), lkind, level, seg))
        start = m.end()
    return out


TITLE_WORDS = ("title", "제목", "titel", "titre", "título", "titolo", "タイトル", "标题", "標題")
SUBTITLE_WORDS = ("subtitle", "부제목", "부제", "untertitel", "sous-titre", "subtítulo", "sottotitolo", "サブタイトル", "副标题")
HEADING_WORDS = ("heading", "머리글", "제목", "überschrift", "titre", "encabezado", "título", "intestazione", "見出し", "标题")


def md_line(sname, lkind, level, txt):
    t = clean_text(txt).strip()
    if not t:
        return ""
    s = (sname or "").lower().strip()
    m = re.match(r"^(\D+?)\s*(\d)$", s)
    if m and any(m.group(1).strip() == w for w in HEADING_WORDS):
        return "#" * min(int(m.group(2)) + 1, 6) + " " + t
    if any(s == w or s == w + " alt" for w in SUBTITLE_WORDS):
        return "## " + t
    if any(s == w or s == w + " alt" for w in TITLE_WORDS):
        return "# " + t
    if s in ("heading", "머리글"):
        return "## " + t
    ind = "  " * int(level)
    if lkind == "number":
        return f"{ind}1. {t}"
    if lkind == "bullet":
        return f"{ind}- {t}"
    return t


def _b64(v):
    import base64
    if isinstance(v, (bytes, bytearray)):
        return bytes(v)
    try:
        return base64.b64decode(v or "")
    except Exception:  # noqa: BLE001
        return b""


def _dec128(b):
    try:
        _patch_numbers_parser()
        from numbers_parser import cell as _nc
        return float(_nc._unpack_decimal128(b))
    except Exception:  # noqa: BLE001
        return None


def _cell_value(buf, strings, rich):
    """Decode one table cell (Pages/Keynote tables use the same storage as Numbers)."""
    import struct
    if not buf or len(buf) < 12:
        return ""
    ver, ctype = buf[0], buf[1]
    flags = struct.unpack("<i", buf[8:12] if ver == 5 else buf[4:8])[0]
    off = 12
    num = dbl = secs = sid = rid = None

    def i4():
        nonlocal off
        v = struct.unpack("<i", buf[off:off + 4])[0]
        off += 4
        return v

    def f8():
        nonlocal off
        v = struct.unpack("<d", buf[off:off + 8])[0]
        off += 8
        return v
    if ver == 5:
        if flags & 0x1:
            num = _dec128(buf[off:off + 16]); off += 16
        if flags & 0x2:
            dbl = f8()
        if flags & 0x4:
            secs = f8()
        if flags & 0x8:
            sid = i4()
        if flags & 0x10:
            rid = i4()
    elif ver == 4:   # pre-2015 storage
        for bit, kind in ((0x2, "i"), (0x80, "i"), (0x400, "i"), (0x800, "i"), (0x4, "i"), (0x8, "i"),
                          (0x10, "s"), (0x20, "d"), (0x40, "t"), (0x200, "r")):
            if flags & bit:
                if kind == "i":
                    i4()
                elif kind == "s":
                    sid = i4()
                elif kind == "r":
                    rid = i4()
                elif kind == "d":
                    dbl = f8()
                else:
                    secs = f8()
    else:
        return ""
    if ctype == 3:
        return strings.get(sid, "")
    if ctype == 9:
        return rich.get(rid, "")
    if ctype in (2, 10):
        v = num if num is not None else dbl
        if v is None:
            return ""
        return str(int(v)) if float(v).is_integer() else f"{v:g}"
    if ctype == 5 and secs is not None:
        return (dt.datetime(2001, 1, 1) + dt.timedelta(seconds=secs)).strftime("%Y-%m-%d")
    if ctype == 6:
        return "TRUE" if (dbl or 0) > 0 else "FALSE"
    if ctype == 7 and dbl is not None:
        return str(dt.timedelta(seconds=dbl))
    if ctype == 8:
        return "#ERROR"
    return ""


def table_grid(iwa, model):
    """Rebuild a Pages/Keynote table as rows of text. Falls back to the plain word list."""
    bds = model.get("baseDataStore", {})
    strings = {e.get("key"): e.get("string", "") for e in iwa.get(bds.get("stringTable")).get("entries", [])}
    rich = {}
    for e in iwa.get(bds.get("richTextTable")).get("entries", []):
        payload = iwa.get(e.get("richTextPayload"))
        st = iwa.get(payload.get("storage"))
        rich[e.get("key")] = clean_text("".join(st.get("text", []))).strip()
    nr, nc = int(model.get("numberOfRows", 0)), int(model.get("numberOfColumns", 0))
    try:
        if not nr or not nc or nr * nc > 200000:
            raise ValueError
        grid = [[""] * nc for _ in range(nr)]
        tsize = bds.get("tiles", {}).get("tileSize") or 256
        for tref in bds.get("tiles", {}).get("tiles", []):
            tile = iwa.get(tref.get("tile"))
            for ri in tile.get("rowInfos", []):
                row = tref.get("tileid", 0) * tsize + ri.get("tileRowIndex", 0)
                if row >= nr:
                    continue
                if ri.get("cellStorageBuffer"):
                    buf, offs, wide = _b64(ri["cellStorageBuffer"]), _b64(ri.get("cellOffsets")), ri.get("hasWideOffsets")
                else:
                    buf, offs, wide = _b64(ri.get("cellStorageBufferPreBnc")), _b64(ri.get("cellOffsetsPreBnc")), False
                import array
                o = array.array("h")
                o.frombytes(offs[: len(offs) // 2 * 2])
                o = [x * 4 if wide else x for x in o.tolist()]
                for col in range(min(nc, len(o))):
                    if o[col] < 0:
                        continue
                    nxt = next((x for x in o[col + 1:] if x >= 0), len(buf))
                    grid[row][col] = _cell_value(buf[o[col]:nxt], strings, rich)
        if any(any(r) for r in grid):
            return grid
    except Exception:  # noqa: BLE001
        pass
    words = [strings[k] for k in sorted(strings)]
    return [words] if words else []


# ───────────────────────── read: Pages ─────────────────────────
def read_pages_iwa(pkg):
    iwa = IWA(pkg)
    lines = []
    body = [(i, o) for i, o in iwa.of_type("TSWP.StorageArchive") if o.get("kind") == "BODY"]
    others = {"HEADER": [], "FOOTNOTE": [], "TEXTBOX": [], "NOTE": []}
    for i, o in iwa.of_type("TSWP.StorageArchive"):
        k = o.get("kind", "TEXTBOX")
        if k in others and o.get("text") and clean_text("".join(o["text"])).strip():
            others[k].append(o)
    for _, st in body:
        for p in storage_paragraphs(iwa, st):
            lines.append(md_line(*p))
    md = "\n\n".join(x for x in lines if x)
    md = re.sub(r"\n{3,}", "\n\n", md)
    extra = []
    tb = [clean_text("".join(o["text"])).strip() for o in others["TEXTBOX"]]
    tb = [t for t in tb if t]
    if tb:
        extra.append("## Text boxes\n\n" + "\n\n".join(tb))
    hf = list(dict.fromkeys(clean_text("".join(o["text"])).strip() for o in others["HEADER"]))
    hf = [t for t in hf if t]
    if hf:
        extra.append("## Headers / footers\n\n" + "\n".join("- " + t for t in hf))
    fn = [clean_text("".join(o["text"])).strip() for o in others["FOOTNOTE"]]
    if fn:
        extra.append("## Footnotes\n\n" + "\n".join(f"{n}. {t}" for n, t in enumerate(fn, 1)))
    tables = [t for _, t in iwa.of_type("TST.TableModelArchive")]
    for n, t in enumerate(tables, 1):
        g = table_grid(iwa, t)
        if g:
            extra.append(f"## Table {n} — {t.get('tableName', '')}\n\n" + md_table(g))
    images = sum(1 for n in pkg.list() if n.startswith("Data/"))
    return {"type": "pages", "markdown": md, "extra": extra, "images": images, "damaged": iwa.failed,
            "paragraphs": sum(1 for x in lines if x), "tables": len(tables)}


# ───────────────────────── read: Keynote ─────────────────────────
def keynote_slides(iwa):
    shows = iwa.of_type("KN.ShowArchive")
    if not shows:
        return []
    nodes = []

    def walk(refs):
        for r in refs or []:
            node = iwa.get(r)
            if node.get("_pbtype") == "KN.SlideNodeArchive":
                nodes.append(node)
                walk(node.get("children"))
    tree = shows[0][1].get("slideTree", {})
    if tree.get("slides"):
        walk(tree["slides"])
    elif tree.get("rootSlideNode"):   # Keynote 6 (2013) keeps one hidden root node
        walk(iwa.get(tree["rootSlideNode"]).get("children"))
    slides = []
    for n, node in enumerate(nodes, 1):
        sid = str(node.get("slide", {}).get("identifier"))
        fname = iwa.where.get(sid, "")
        texts, notes = [], []
        for i in iwa.order:
            if iwa.where[i] != fname:
                continue
            o = iwa.objs[i]
            if o.get("_pbtype") != "TSWP.StorageArchive":
                continue
            paras = [md_line("", ln, lv, t) for _, ln, lv, t in storage_paragraphs(iwa, o)]
            block = "\n".join(p for p in paras if p)
            if not block.strip():
                continue
            (notes if o.get("kind") == "NOTE" else texts).append(block)
        thumb = None
        th = node.get("thumbnails") or []
        if th:
            thumb = iwa.data_file(th[0].get("identifier"))
        slides.append({"n": n, "skipped": bool(node.get("isSkipped")), "texts": texts,
                       "notes": notes, "thumb": thumb, "file": fname})
    return slides


def read_keynote_iwa(pkg):
    iwa = IWA(pkg)
    slides = keynote_slides(iwa)
    parts = []
    for s in slides:
        head = f"## Slide {s['n']}" + (" (skipped)" if s["skipped"] else "")
        body = "\n\n".join(s["texts"]) or "_(no text — picture or shapes only)_"
        note = ("\n\n> **Speaker notes:** " + " / ".join(x.replace("\n", " ") for x in s["notes"])) if s["notes"] else ""
        parts.append(f"{head}\n\n{body}{note}")
    shows = iwa.of_type("KN.ShowArchive")
    size = shows[0][1].get("size", {}) if shows else {}
    return {"type": "keynote", "markdown": "\n\n".join(parts), "extra": [], "slides": slides, "damaged": iwa.failed,
            "size": (size.get("width", 1024), size.get("height", 768)),
            "images": sum(1 for n in pkg.list() if n.startswith("Data/"))}


# ───────────────────────── read: Numbers ─────────────────────────
_PATCHED = False


def _patch_numbers_parser():
    """numbers-parser stores numbers through float maths, so 27 becomes 26.999999999999996.
    Replace its decimal128 packing with exact Decimal maths (same layout Numbers uses)."""
    global _PATCHED
    if _PATCHED:
        return
    from decimal import Decimal, getcontext
    import numbers_parser.cell as nc
    from numbers_parser.constants import DECIMAL128_BIAS
    getcontext().prec = 40

    def pack(value):
        buf = bytearray(16)
        d = Decimal(repr(float(value))) if not isinstance(value, Decimal) else value
        if d == 0:
            exp_u = -16
            mant = 0
        else:
            exp_u = d.copy_abs().adjusted() - 16
            mant = int((d.copy_abs().scaleb(-exp_u)).to_integral_value())
        exp = exp_u + DECIMAL128_BIAS
        buf[15] |= exp >> 7
        buf[14] |= (exp & 0x7F) << 1
        i = 0
        while mant >= 1 and i < 14:
            buf[i] = mant & 0xFF
            mant >>= 8
            i += 1
        if d < 0:
            buf[15] |= 0x80
        return buf

    def unpack(buf):
        exp = (((buf[15] & 0x7F) << 7) | (buf[14] >> 1)) - DECIMAL128_BIAS
        mant = buf[14] & 1
        for i in range(13, -1, -1):
            mant = mant * 256 + buf[i]
        v = Decimal(mant).scaleb(exp)
        return float(-v if buf[15] & 0x80 else v)
    nc._pack_decimal128 = pack
    nc._unpack_decimal128 = unpack
    _PATCHED = True


class NumbersOpenError(Exception):
    pass


def open_numbers(path, password=None, soft=False):
    need("numbers_parser", "numbers-parser")
    from numbers_parser import Document
    _patch_numbers_parser()
    try:
        return Document(str(path), password=password) if password else Document(str(path))
    except Exception as e:  # noqa: BLE001
        msg = str(e)
        hint = ""
        try:
            pkg = Package(path)
            for n in pkg.list():
                if n.endswith(".iwph"):
                    hint = pkg.read(n).decode("utf-8", "ignore").strip()
        except SystemExit:
            pass
        hint_txt = f" (password hint saved in the file: {hint!r})" if hint else ""
        if "password" in msg.lower() and "No password" in msg:
            die("this Numbers file is password-protected — run again with --password PW" + hint_txt)
        if "password" in msg.lower() or "decrypt" in msg.lower():
            die("wrong password for this Numbers file" + hint_txt)
        if soft:
            raise NumbersOpenError(msg)
        die(f"could not open this Numbers file ({type(e).__name__}: {msg}). It may be damaged. "
            "Older Numbers files (before 2015) can be read and converted but not edited.")


def cell_text(c):
    v = c.value
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        f = getattr(c, "formatted_value", "") or ""
        if not f or re.fullmatch(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}", f):
            return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
        return f
    if isinstance(v, float):
        f = getattr(c, "formatted_value", "") or ""
        if not f or re.fullmatch(r"-?\d+(\.\d+)?(e[-+]?\d+)?", f):
            return str(int(v)) if v.is_integer() and abs(v) < 1e15 else f"{v:.15g}"
        return f
    try:
        f = c.formatted_value
        if f not in (None, ""):
            return str(f)
    except Exception:  # noqa: BLE001
        pass
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def md_table(rows):
    rows = [[(x or "").replace("|", "\\|").replace("\n", " ") for x in r] for r in rows]
    while rows and not any(rows[-1]):
        rows.pop()
    if not rows:
        return "_(empty table)_"
    w = max(len(r) for r in rows)
    while w > 1 and not any(len(r) >= w and r[w - 1] for r in rows):
        w -= 1
    rows = [(r + [""] * w)[:w] for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * w]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def read_numbers_iwa(pkg):
    iwa = IWA(pkg)
    parts, grids = [], []
    for sid, sh in iwa.of_type("TN.SheetArchive"):
        sheet = sh.get("name", "Sheet")
        parts.append(f"## Sheet: {sheet}")
        for ref in sh.get("drawableInfos", []):
            info = iwa.get(ref)
            model = iwa.get(info.get("tableModel")) if info.get("_pbtype") == "TST.TableInfoArchive" else {}
            if not model:
                continue
            rows = table_grid(iwa, model)
            name = model.get("tableName", "Table")
            parts.append(f"### Table: {name}\n\n{md_table(rows)}")
            grids.append((sheet, name, rows))
    if not grids:   # sheet list not found: show every table
        for _, model in iwa.of_type("TST.TableModelArchive"):
            rows = table_grid(iwa, model)
            parts.append(f"### Table: {model.get('tableName', 'Table')}\n\n{md_table(rows)}")
            grids.append(("Sheet", model.get("tableName", "Table"), rows))
    return {"type": "numbers", "markdown": "\n\n".join(parts), "extra": [], "old": True, "grids": grids,
            "tables": len(grids), "damaged": iwa.failed}


def read_numbers(path, password=None, formulas=False, soft=False):
    doc = open_numbers(path, password, soft=soft)
    parts, n_tables, n_formulas = [], 0, 0
    for s in doc.sheets:
        parts.append(f"## Sheet: {s.name}")
        for t in s.tables:
            n_tables += 1
            rows = []
            for r in t.rows():
                row = []
                for c in r:
                    txt = cell_text(c)
                    if c.formula:
                        n_formulas += 1
                        if formulas:
                            txt = f"{txt} `={c.formula}`"
                    row.append(txt)
                rows.append(row)
            parts.append(f"### Table: {t.name} ({t.num_rows} rows × {t.num_cols} columns)\n\n{md_table(rows)}")
    return {"type": "numbers", "markdown": "\n\n".join(parts), "extra": [], "doc": doc,
            "tables": n_tables, "formulas": n_formulas}


# ───────────────────────── read: old XML formats (iWork '09 and earlier) ─────────────────────────
def read_xml(pkg):
    names = pkg.list()
    for cand in ("index.xml", "index.xml.gz", "index.apxl", "index.apxl.gz", "presentation.apxl",
                 "presentation.apxl.gz"):
        if cand in names:
            data = pkg.read(cand)
            if cand.endswith(".gz"):
                data = gzip.decompress(data)
            break
    else:
        die("no readable content found")
    from xml.etree import ElementTree as ET
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        # some old files have broken namespaces — fall back to pulling paragraph text out by pattern
        txt = data.decode("utf-8", errors="ignore")
        paras = [re.sub(r"<[^>]+>", "", m).strip() for m in re.findall(r"<sf:p\b[^>]*>(.*?)</sf:p>", txt, flags=re.S)]
        import html as _h
        paras = [_h.unescape(p) for p in paras if p.strip()]
        return {"type": pkg.kind, "markdown": "\n\n".join(paras), "extra": [], "old": True}

    def local(tag):
        return tag.rsplit("}", 1)[-1]
    parts, grids = [], []
    if pkg.kind == "keynote":
        for i, slide in enumerate((e for e in root.iter() if local(e.tag) == "slide"), 1):
            paras = ["".join(p.itertext()).strip() for p in slide.iter() if local(p.tag) == "p"]
            parts.append(f"## Slide {i}\n\n" + "\n\n".join(x for x in paras if x))
    elif pkg.kind == "numbers":
        def attr(e, name):
            return next((v for k, v in e.attrib.items() if local(k) == name), None)

        def num(v):
            try:
                f = float(v)
                return str(int(f)) if f.is_integer() else f"{round(f, 10):g}"
            except (TypeError, ValueError):
                return v or ""

        def val(c):
            tag = local(c.tag)
            if tag == "t":
                ct = next((x for x in c.iter() if local(x.tag) == "ct"), None)
                return (attr(ct, "s") if ct is not None else None) or "".join(c.itertext()).strip()
            if tag == "n":
                return num(attr(c, "v"))
            if tag == "d":
                secs = attr(c, "cell-date")
                return (dt.datetime(2001, 1, 1) + dt.timedelta(seconds=float(secs))).strftime("%Y-%m-%d") if secs else ""
            if tag == "b":
                return "TRUE" if attr(c, "v") in ("1", "true") else "FALSE"
            if tag == "du":
                return attr(c, "du") or attr(c, "v") or ""
            if tag == "pm":   # pop-up menu cell: show the chosen item
                ref = next((attr(x, "IDREF") for x in c.iter() if local(x.tag) == "proxied-cell-ref"), None)
                for x in c.iter():
                    if ref and attr(x, "ID") == ref:
                        return val(x)
                return ""
            if tag == "f":
                for x in c.iter():
                    lt = local(x.tag)
                    if lt == "rn":
                        return num(attr(x, "v"))
                    if lt in ("rt", "rs"):
                        ct = next((y for y in x.iter() if local(y.tag) == "ct"), None)
                        return (attr(ct, "s") if ct is not None else "") or "".join(x.itertext()).strip()
                    if lt == "rb":
                        return "TRUE" if attr(x, "v") in ("1", "true") else "FALSE"
                    if lt == "rd":
                        secs = attr(x, "cell-date")
                        return (dt.datetime(2001, 1, 1) + dt.timedelta(seconds=float(secs))).strftime("%Y-%m-%d") if secs else ""
                return ""
            return ""
        for tab in (e for e in root.iter() if local(e.tag) == "tabular-model"):
            name = attr(tab, "name") or "Table"
            grid = next((e for e in tab.iter() if local(e.tag) == "grid"), None)
            ncols = int(attr(grid, "numcols") or 1) if grid is not None else 1
            ds = next((e for e in tab.iter() if local(e.tag) == "datasource"), None)
            cells = []
            for c in (list(ds) if ds is not None else []):
                if local(c.tag) not in ("t", "n", "d", "f", "b", "du", "g", "s", "o", "e", "pm"):
                    continue
                cells.append(val(c))
                span = int(attr(c, "col-span") or 1)
                cells += [""] * (span - 1)
            rows = [cells[i:i + ncols] for i in range(0, len(cells), ncols)]
            parts.append(f"### Table: {name}\n\n" + md_table(rows))
            grids.append(("Sheet", name, rows))
    else:
        body = next((e for e in root.iter() if local(e.tag) == "text-body"), root)
        paras = ["".join(p.itertext()).strip() for p in body.iter() if local(p.tag) == "p"]
        parts.append("\n\n".join(x for x in paras if x))
    md = "\n\n".join(parts)
    if not re.sub(r"#+ Slide \d+|\s", "", md):
        md += "\n\n_(no text found — this very old file may hold only pictures)_"
    return {"type": pkg.kind, "markdown": md, "extra": [], "old": True, "grids": grids}


# ───────────────────────── dispatch ─────────────────────────
def load(path, password=None, formulas=False):
    pkg = Package(path)
    if pkg.kind is None:
        die("not an iWork file — expected .pages, .numbers or .key")
    fmt = pkg.format
    if fmt == "locked" and pkg.kind != "numbers":
        die("this file is password-protected. Ask the owner to remove the password "
            "(File > Set Password) or to export it as Word/PowerPoint/PDF.")
    if fmt == "xml":
        return read_xml(pkg)
    if fmt == "unknown" and pkg.kind != "numbers":
        die("unrecognised iWork package (very old or damaged file)")
    if pkg.kind == "numbers":
        try:
            return read_numbers(path, password, formulas, soft=fmt == "iwa")
        except NumbersOpenError:
            return read_numbers_iwa(pkg)   # 2013-2014 Numbers files: read through the shared table decoder
    if pkg.kind == "pages":
        return read_pages_iwa(pkg)
    return read_keynote_iwa(pkg)


def full_markdown(r, title):
    md = f"# {title}\n\n"
    if r.get("damaged"):
        md += f"> **Warning:** {len(r['damaged'])} part(s) of this file are damaged and were skipped: " \
              + ", ".join(r["damaged"][:5]) + "\n\n"
    md += r["markdown"]
    if r.get("extra"):
        md += "\n\n" + "\n\n".join(r["extra"])
    return md.strip() + "\n"


# ───────────────────────── commands ─────────────────────────
def cmd_info(a):
    pkg = Package(a.file)
    size = sum(1 for _ in pkg.list())
    info = {"file": Path(a.file).name, "type": pkg.kind, "storage": "folder bundle" if pkg.path.is_dir() else "zip",
            "format": {"iwa": "modern (2013+)", "xml": "old iWork '09 or earlier", "locked": "password-protected",
                       "unknown": "unknown"}[pkg.format], "parts": size}
    if pkg.has("Metadata/BuildVersionHistory.plist"):
        m = re.findall(rb"<string>([^<]+)</string>", pkg.read("Metadata/BuildVersionHistory.plist"))
        if m:
            info["saved_by"] = m[-1].decode(errors="ignore")
    if pkg.format == "locked" and pkg.kind != "numbers":
        info["note"] = "password-protected: cannot be read without removing the password on a Mac"
    elif pkg.kind == "numbers" and pkg.format != "xml" and (pkg.format != "locked" or a.password):
        try:
            doc = open_numbers(a.file, a.password, soft=True)
        except NumbersOpenError:
            r = read_numbers_iwa(pkg)
            info["tables"] = [{"sheet": g[0], "name": g[1], "rows": len(g[2])} for g in r["grids"]]
            info["note"] = "older Numbers file (2013-2014): can be read and converted, not edited"
            print(json.dumps(info, ensure_ascii=False, indent=2))
            return
        info["sheets"] = [{"name": s.name, "tables": [{"name": t.name, "rows": t.num_rows, "cols": t.num_cols,
                                                         "merged": len(t.merge_ranges)} for t in s.tables]}
                          for s in doc.sheets]
        info["formulas"] = sum(1 for s in doc.sheets for t in s.tables for r in t.rows() for c in r if c.formula)
        info["editable"] = True
    elif pkg.format != "locked":
        r = load(a.file)
        if r["type"] == "keynote" and "slides" in r:
            info["slides"] = len(r["slides"])
            info["slides_with_notes"] = sum(1 for s in r["slides"] if s["notes"])
            info["skipped_slides"] = sum(1 for s in r["slides"] if s["skipped"])
        if r["type"] == "pages":
            info["paragraphs"] = r.get("paragraphs")
            info["tables"] = r.get("tables")
        info["words"] = len(re.findall(r"\w+", r["markdown"]))
        info["editable"] = "text find-and-replace" if pkg.format == "iwa" else "no (old format: read/convert only)"
        info["media_files"] = r.get("images", 0)
    if pkg.format == "locked" and pkg.kind == "numbers" and not a.password:
        info["note"] = "password-protected Numbers file: run again with --password PW"
    print(json.dumps(info, ensure_ascii=False, indent=2))


def cmd_read(a):
    r = load(a.file, a.password, a.formulas)
    md = full_markdown(r, Path(a.file).name)
    if a.out:
        Path(a.out).write_text(md, encoding="utf-8")
        print(f"saved: {a.out} ({len(md)} chars)")
    else:
        print(md)


def cmd_preview(a):
    pkg = Package(a.file)
    out = Path(a.out or (Path(a.file).stem + "_preview"))
    out.mkdir(parents=True, exist_ok=True)
    names = pkg.list()
    saved = []
    for cand in ("preview.jpg", "QuickLook/Preview.pdf", "QuickLook/Thumbnail.jpg", "preview-web.jpg"):
        if cand in names:
            dest = out / ("cover" + Path(cand).suffix)
            dest.write_bytes(pkg.read(cand))
            saved.append(str(dest))
            break
    if pkg.kind == "keynote" and pkg.format == "iwa":
        iwa = IWA(pkg)
        for s in keynote_slides(iwa):
            if s["thumb"]:
                dest = out / f"slide_{s['n']:03d}{Path(s['thumb']).suffix}"
                dest.write_bytes(pkg.read(s["thumb"]))
                saved.append(str(dest))
    if not saved:
        die("this file has no built-in preview picture")
    print("saved previews (small pictures made by the Mac when the file was saved):")
    for s in saved:
        print("  " + s)


def cmd_images(a):
    pkg = Package(a.file)
    out = Path(a.out or (Path(a.file).stem + "_images"))
    names = pkg.list()
    media = [n for n in names if n.startswith("Data/") and not n.endswith("/")]
    if pkg.format == "iwa":
        iwa = IWA(pkg)
        if not a.all:
            if pkg.kind == "keynote":
                keep_files = {s["file"] for s in keynote_slides(iwa)}
            else:
                keep_files = {f for f in iwa.datarefs if "Master" not in f and "Theme" not in f}
            used = set()
            for f in keep_files:
                for ident in iwa.datarefs.get(f, ()):
                    df = iwa.data_file(ident)
                    if df:
                        used.add(df)
            media = [m for m in media if m in used]
        nice = {}
        for _, o in iwa.of_type("TSP.PackageMetadata"):
            for x in o.get("datas", []):
                df = iwa.data_file(x.get("identifier"))
                if df and x.get("preferredFileName"):
                    nice[df] = x["preferredFileName"]
    else:
        nice = {}
        if not a.all:
            media = [m for m in media if not Path(m).name.startswith(("st", "mt"))]
    media = [m for m in media if not re.match(r"Data/(st|mt)-", m) or a.all]
    if not media:
        print("no pictures or media used in this document" + ("" if a.all else " (try --all to include theme pictures)"))
        return
    out.mkdir(parents=True, exist_ok=True)
    seen = set()
    for m in media:
        name = re.sub(r'[\\/:*?"<>|]', "_", nice.get(m) or Path(m).name)
        stem, ext = os.path.splitext(name)
        k = 2
        while name.lower() in seen:
            name = f"{stem}-{k}{ext}"
            k += 1
        seen.add(name.lower())
        (out / name).write_bytes(pkg.read_any(m))
        print(f"  {out / name}  ({len(pkg.read_any(m)) // 1024} KB)")
    print(f"saved {len(media)} file(s) to {out}")


# ── convert ──
EXCEL_FUNCS = set("""ABS ACOS ACOSH ADDRESS AND AREAS ASIN ASINH ATAN ATAN2 ATANH AVEDEV AVERAGE AVERAGEA AVERAGEIF
AVERAGEIFS BESSELI BESSELJ BESSELK BESSELY BIN2DEC BIN2HEX BIN2OCT BINOMDIST CEILING CHAR CHIDIST CHIINV CHITEST CHOOSE
CLEAN CODE COLUMN COLUMNS COMBIN CONCATENATE CONFIDENCE CONVERT CORREL COS COSH COUNT COUNTA COUNTBLANK COUNTIF
COUNTIFS COUPDAYBS COUPDAYS COUPDAYSNC COUPNUM COVAR CRITBINOM CUMIPMT CUMPRINC DATE DATEDIF DATEVALUE DAY DAYS360 DB DDB
DEC2BIN DEC2HEX DEC2OCT DEGREES DELTA DEVSQ DISC DOLLAR EDATE EFFECT EOMONTH ERF ERFC EVEN EXACT EXP EXPONDIST FACT
FACTDOUBLE FALSE FDIST FIND FINV FISHER FISHERINV FIXED FLOOR FORECAST FREQUENCY FV FVSCHEDULE GAMMADIST GAMMAINV
GAMMALN GCD GEOMEAN GESTEP GROWTH HARMEAN HEX2BIN HEX2DEC HEX2OCT HLOOKUP HOUR HYPERLINK IF IFERROR INDEX INDIRECT INT
INTERCEPT IPMT IRR ISBLANK ISERROR ISEVEN ISNUMBER ISODD ISPMT ISTEXT LARGE LCM LEFT LEN LINEST LN LOG LOG10 LOGINV
LOGNORMDIST LOOKUP LOWER MATCH MAX MAXA MEDIAN MID MIN MINA MINUTE MIRR MOD MODE MONTH MROUND MULTINOMIAL N NETWORKDAYS NOMINAL
NORMDIST NORMINV NORMSDIST NORMSINV NOT NOW NPER NPV OCT2BIN OCT2DEC OCT2HEX ODD OFFSET OR PEARSON PERCENTILE PERCENTRANK
PERMUT PI PMT POISSON POWER PPMT PRICE PRICEDISC PRICEMAT PROB PRODUCT PROPER PV QUARTILE QUOTIENT RADIANS RAND
RANDBETWEEN RANK RATE RECEIVED REPLACE REPT RIGHT ROMAN ROUND ROUNDDOWN ROUNDUP ROW ROWS RSQ SEARCH SECOND SERIESSUM SIGN
SIN SINH SLOPE SMALL SQRT SQRTPI STANDARDIZE STDEV STDEVA STDEVP STDEVPA SUBSTITUTE SUBTOTAL SUM SUMIF SUMIFS SUMPRODUCT SUMSQ
SUMX2MY2 SUMX2PY2 SUMXMY2 T TAN TANH TDIST TEXT TIME TIMEVALUE TINV TODAY TRANSPOSE TREND TRIM TRUE TRUNC TTEST TYPE UPPER
VALUE VAR VARA VARP VARPA VDB VLOOKUP WEEKDAY WEEKNUM WORKDAY XIRR XNPV YEAR YEARFRAC YIELD YIELDDISC YIELDMAT ZTEST""".split())
EXCEL_NEW = {"XLOOKUP", "XMATCH", "IFS", "MAXIFS", "MINIFS", "CONCAT", "TEXTJOIN", "SWITCH", "IFNA", "DAYS",
             "ISOWEEKNUM", "FILTER", "SORT", "UNIQUE", "SEQUENCE", "TEXTBEFORE", "TEXTAFTER", "STDEV.S", "STDEV.P",
             "VAR.S", "VAR.P", "CEILING.MATH", "FLOOR.MATH", "BITAND", "BITOR", "BITXOR", "ISFORMULA", "FORMULATEXT"}


def numbers_formula_to_excel(f, sheet_of, here):
    """Translate a Numbers formula to Excel. Returns None when it can't be done safely.
    sheet_of(sheet, table) -> Excel sheet name or None; here = (sheet, table) of the cell."""
    parts = re.split(r'("(?:[^"]|"")*")', f)
    out = []
    for i, p in enumerate(parts):
        if i % 2:   # inside a string literal
            out.append(p)
            continue
        p = p.replace("×", "*").replace("÷", "/").replace("≠", "<>").replace("≤", "<=").replace("≥", ">=")

        def qual(m):
            names = [x.strip() for x in m.group(1).split("::")]
            ref = m.group(2)
            if len(names) == 2:
                sh, tb = names
            else:
                sh, tb = here[0], names[0]
            target = sheet_of(sh, tb)
            if target is None:
                raise ValueError("unknown table")
            if re.fullmatch(r"\$?[A-Z]{1,3}", ref):   # whole column
                ref = f"{ref}:{ref}"
            elif re.fullmatch(r"\$?[A-Z]{1,3}:\$?[A-Z]{1,3}", ref) or re.fullmatch(r"\$?\d+(:\$?\d+)?", ref):
                pass
            if (sh, tb) == here:
                return ref
            return "'" + target.replace("'", "''") + "'!" + ref
        try:
            p = re.sub(r"((?:[^():,+\-*/&=<>^%\s\"]+(?:\s[^():,+\-*/&=<>^%\s\"]+)*::){1,2})"
                       r"(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}(?::\$?[A-Z]{1,3})?)", qual, p)
        except ValueError:
            return None
        if "::" in p:
            return None
        # every remaining word must be a known Excel function, TRUE/FALSE, a cell/range or a sheet reference
        scrub = re.sub(r"'(?:[^']|'')*'!", "", p)
        for m in re.finditer(r"[A-Za-z_][A-Za-z0-9_.]*", scrub):
            w = m.group(0)
            nxt = scrub[m.end():m.end() + 1]
            if nxt == "(":
                if w.upper() in EXCEL_FUNCS:
                    continue
                if w.upper() in EXCEL_NEW:
                    continue
                return None
            if w.upper() in ("TRUE", "FALSE"):
                continue
            if re.fullmatch(r"[A-Z]{1,3}\d+", w) or re.fullmatch(r"[A-Z]{1,3}", w):
                continue
            return None
        for fn in EXCEL_NEW:   # newer functions need Excel's _xlfn. prefix in the file
            p = re.sub(rf"(?<![\w.]){re.escape(fn)}\(", f"_xlfn.{fn}(", p)
        out.append(p)
    return "=" + "".join(out)


def numbers_to_xlsx(r, out):
    need("openpyxl", "openpyxl")
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    doc = r["doc"]
    wb = Workbook()
    wb.remove(wb.active)
    used, names = set(), {}
    for s in doc.sheets:
        for t in s.tables:
            base = s.name if len(s.tables) == 1 else f"{s.name} - {t.name}"
            name = re.sub(r"[\\/*?:\[\]]", "_", base)[:31] or "Sheet"
            k = 2
            while name.lower() in used:
                name = f"{base[:27]} ({k})"
                k += 1
            used.add(name.lower())
            names[(s.name, t.name)] = name
    stats = {"kept": 0, "value": 0}

    def sheet_of(sh, tb):
        if (sh, tb) in names:
            return names[(sh, tb)]
        hits = [v for (a_, b_), v in names.items() if b_ == tb]
        return hits[0] if len(hits) == 1 else None
    for s in doc.sheets:
        for t in s.tables:
            ws = wb.create_sheet(names[(s.name, t.name)])
            for ri, row in enumerate(t.rows()):
                for ci, c in enumerate(row):
                    v = c.value
                    if v is None or type(c).__name__ == "MergedCell":
                        continue
                    if isinstance(v, dt.timedelta):
                        v = cell_text(c)
                    elif isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
                        v = int(v)
                    xf = None
                    if c.formula and r.get("live_formulas", True):
                        try:
                            xf = numbers_formula_to_excel(c.formula, sheet_of, (s.name, t.name))
                        except Exception:  # noqa: BLE001
                            xf = None
                    cell = ws.cell(row=ri + 1, column=ci + 1, value=xf or v)
                    if isinstance(c.value, dt.datetime):   # US-style date display
                        cell.number_format = "mm/dd/yyyy" if (c.value.hour, c.value.minute) == (0, 0) else "mm/dd/yyyy h:mm AM/PM"
                    if c.formula:
                        if xf:
                            stats["kept"] += 1
                        else:
                            stats["value"] += 1
                            cell.comment = Comment(f"Numbers formula (kept as value): ={c.formula}", "iwork-file")
                    try:
                        st = c.style
                        if st.bold or st.italic:
                            cell.font = Font(bold=st.bold, italic=st.italic)
                        bg = st.bg_color
                        if bg is not None and not isinstance(bg, list):
                            hexc = f"{bg.r:02X}{bg.g:02X}{bg.b:02X}"
                            if hexc != "FFFFFF":
                                cell.fill = PatternFill("solid", fgColor=hexc)
                    except Exception:  # noqa: BLE001
                        pass
            for rng in t.merge_ranges:
                try:
                    ws.merge_cells(rng)
                except Exception:  # noqa: BLE001
                    pass
            for ci in range(t.num_cols):
                try:
                    ws.column_dimensions[get_column_letter(ci + 1)].width = max(8, min(60, t.col_width(ci) / 7))
                except Exception:  # noqa: BLE001
                    pass
    if stats["kept"]:
        wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return stats


def grids_to_xlsx(grids, out):
    need("openpyxl", "openpyxl")
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for sh, name, rows in grids:
        base = re.sub(r"[\\/*?:\[\]]", "_", name if sh == "Sheet" else f"{sh} - {name}")[:31] or "Sheet"
        nm, k = base, 2
        while nm.lower() in used:
            nm = f"{base[:27]} ({k})"
            k += 1
        used.add(nm.lower())
        ws = wb.create_sheet(nm)
        for r in rows:
            out_row = []
            for v in r:
                if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v or ""):
                    out_row.append(dt.datetime.strptime(v, "%Y-%m-%d"))
                    continue
                try:
                    out_row.append(float(v) if re.fullmatch(r"-?\d+(\.\d+)?", v or "") else v)
                except ValueError:
                    out_row.append(v)
            ws.append([int(x) if isinstance(x, float) and x.is_integer() else x for x in out_row])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, dt.datetime):
                    cell.number_format = "mm/dd/yyyy"
    if not wb.worksheets:
        wb.create_sheet("Sheet")
    wb.save(out)


def numbers_to_csv(r, out):
    doc = r["doc"]
    tables = [(s, t) for s in doc.sheets for t in s.tables]
    outp = Path(out)
    written = []
    for i, (s, t) in enumerate(tables):
        p = outp if len(tables) == 1 else outp.with_name(
            f"{outp.stem}_{re.sub(r'[^0-9A-Za-z가-힣_-]+', '_', s.name)}_{re.sub(r'[^0-9A-Za-z가-힣_-]+', '_', t.name)}.csv")
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            for row in t.rows():
                w.writerow([cell_text(c) for c in row])
        written.append(str(p))
    return written


def md_to_docx(md, out, title):
    need("docx", "python-docx")
    from docx import Document as Docx
    d = Docx()
    d.core_properties.title = title
    for block in md.split("\n\n"):
        b = block.strip()
        if not b:
            continue
        if b == "---":
            d.add_page_break()
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", b)
        if m:
            d.add_heading(m.group(2), level=min(len(m.group(1)) - 1, 4) or 0)
            continue
        if b.startswith("|") and "\n|" in b:
            rows = [[x.strip() for x in ln.strip().strip("|").split("|")] for ln in b.splitlines()
                    if not re.match(r"^\|(\s*-+\s*\|)+$", ln.strip())]
            w = max(len(r) for r in rows)
            tb = d.add_table(rows=len(rows), cols=w)
            tb.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    tb.cell(i, j).text = v
            continue
        for ln in b.splitlines():
            lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
            if lm:
                style = "List Number" if lm.group(2)[0].isdigit() else "List Bullet"
                d.add_paragraph(lm.group(3), style=style)
            elif ln.startswith("> "):
                d.add_paragraph(ln[2:].replace("**", ""), style="Intense Quote")
            else:
                d.add_paragraph(ln)
    d.save(out)


def keynote_to_pptx(r, out):
    need("pptx", "python-pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    w, h = r.get("size") or (1024, 768)
    wide = w / max(h, 1) > 1.5
    factor = 13.333 / 10 if wide else 1.0
    if wide:   # 16:9 deck: widen the page and stretch the template boxes to match
        prs.slide_width = Inches(13.333)
    for s in r["slides"]:
        texts = [t for t in s["texts"] if t.strip()]
        layout = prs.slide_layouts[1] if len(texts) > 1 else prs.slide_layouts[5 if not texts else 0]
        sl = prs.slides.add_slide(layout)
        if factor != 1.0:
            for ph in sl.placeholders:
                l, t, w_, h_ = ph.left, ph.top, ph.width, ph.height
                ph.left, ph.top, ph.width, ph.height = int(l * factor), t, int(w_ * factor), h_
        if texts and sl.shapes.title is not None:
            sl.shapes.title.text = re.sub(r"^#+\s*|^-\s|^\d+\.\s", "", texts[0].splitlines()[0])
            rest = texts[0].splitlines()[1:] + [ln for t in texts[1:] for ln in t.splitlines()]
            if layout == prs.slide_layouts[0] and len(sl.placeholders) > 1:
                sl.placeholders[1].text = "\n".join(rest)
            elif rest:
                body = sl.placeholders[1].text_frame
                body.clear()
                first = True
                for ln in rest:
                    m = re.match(r"^(\s*)(?:-|\d+\.)\s+(.*)", ln)
                    p = body.paragraphs[0] if first else body.add_paragraph()
                    p.text = m.group(2) if m else ln
                    p.level = min(len(m.group(1)) // 2, 4) if m else 0
                    first = False
        if s["notes"]:
            sl.notes_slide.notes_text_frame.text = "\n".join(s["notes"])
        if not texts:
            tb = sl.shapes.add_textbox(Inches(1), Inches(3), int(prs.slide_width - Inches(2)), Inches(1))
            tb.text_frame.text = f"(Slide {s['n']}: picture or shapes only)"
            tb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    prs.save(out)


def cmd_convert(a):
    src, dst = Path(a.file), Path(a.out)
    ext = dst.suffix.lower()
    r = load(src, a.password)
    t = r["type"]
    if ext == ".pdf":
        pkg = Package(src)
        if "QuickLook/Preview.pdf" in pkg.list():
            dst.write_bytes(pkg.read("QuickLook/Preview.pdf"))
            print(f"saved: {dst}  (the full PDF the Mac stored inside this file)")
            return
        die("this file has no PDF inside. Convert it to .docx/.pptx/.xlsx and print that to PDF, "
            "or ask the sender to use File › Export To › PDF on their Mac")
    if ext in (".md", ".txt"):
        md = full_markdown(r, src.name)
        if ext == ".txt":
            md = re.sub(r"^#+\s*", "", md, flags=re.M).replace("**", "").replace("`", "")
        dst.write_text(md, encoding="utf-8")
        print(f"saved: {dst}")
        return
    if t == "numbers" and ext == ".xlsx":
        if r.get("old"):
            grids_to_xlsx(r["grids"], dst)
            print(f"saved: {dst}  (older Numbers file: values only)")
            return
        r["live_formulas"] = not a.values_only
        st = numbers_to_xlsx(r, dst)
        msg = f"saved: {dst}"
        if st["kept"] or st["value"]:
            msg += (f"  ({st['kept']} formula(s) work in Excel; {st['value']} Numbers-only formula(s) saved as "
                    "their value with the original formula in a cell note)")
        print(msg)
        return
    if t == "numbers" and ext == ".csv":
        if r.get("old"):
            g = r["grids"]
            for i, (sh, name, rows) in enumerate(g):
                p = dst if len(g) == 1 else dst.with_name(f"{dst.stem}_{i + 1}_{re.sub(r'[^0-9A-Za-z가-힣_-]+', '_', name)}.csv")
                with open(p, "w", newline="", encoding="utf-8-sig") as f:
                    csv.writer(f).writerows(rows)
                print(f"saved: {p}")
            return
        for p in numbers_to_csv(r, dst):
            print(f"saved: {p}")
        return
    if t == "pages" and ext == ".docx":
        md_to_docx(full_markdown(r, src.stem), dst, src.stem)
        print(f"saved: {dst}  (text, headings, lists and tables; pictures and exact layout are not copied)")
        return
    if t == "keynote" and ext == ".pptx":
        if "slides" not in r:   # old Keynote '09: rebuild slides from the text
            r["slides"] = []
            for n, chunk in enumerate(re.split(r"^## Slide \d+\s*$", r["markdown"], flags=re.M)[1:], 1):
                texts = [b.strip() for b in chunk.strip().split("\n\n") if b.strip() and not b.startswith("_(")]
                r["slides"].append({"n": n, "texts": texts, "notes": [], "skipped": False})
        keynote_to_pptx(r, dst)
        print(f"saved: {dst}  (one slide per slide with its text and speaker notes; design is not copied)")
        return
    if t == "keynote" and ext == ".docx":
        md_to_docx(full_markdown(r, src.stem), dst, src.stem)
        print(f"saved: {dst}")
        return
    die(f"cannot convert {t} to {ext}. Numbers → .xlsx .csv .md · Pages → .docx .md .txt · Keynote → .pptx .docx .md .txt")


# ── Numbers editing ──
def parse_value(s):
    if s.startswith("="):
        die("writing formulas into Numbers files is not supported yet — write the value instead")
    t = s.strip()
    m = re.fullmatch(r"(-)?\$?(-)?(\d{1,3}(?:,\d{3})+|\d+)?(\.\d+)?", t)   # 1200, 1,200.50, $1,200, -$5
    zip_like = bool(re.fullmatch(r"0\d+", t))   # keep ZIP codes / IDs like 02134 as text
    if m and (m.group(3) or m.group(4)) and t not in ("-", "$") and not zip_like:
        num = (m.group(3) or "0").replace(",", "") + (m.group(4) or "")
        sign = -1 if (m.group(1) or m.group(2)) else 1
        return sign * (float(num) if "." in num else int(num))
    m = re.fullmatch(r"(-?\d+(?:\.\d+)?)%", t)
    if m:
        return float(m.group(1)) / 100
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", t):
        return dt.datetime.strptime(t, "%Y-%m-%d")
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)   # US: month/day/year
    if m:
        return dt.datetime(int(m.group(3)), int(m.group(1)), int(m.group(2)))
    if s.lower() in ("true", "false"):
        return s.lower() == "true"
    if s == "":
        return ""
    return s


def pick_table(doc, sheet, table):
    sheets = doc.sheets
    if sheet:
        sheets = [s for s in sheets if s.name == sheet]
        if not sheets:
            die(f"no sheet named {sheet!r}. Sheets: {[s.name for s in doc.sheets]}")
    s = sheets[0]
    ts = s.tables
    if table:
        ts = [t for t in ts if t.name == table]
        if not ts:
            die(f"no table named {table!r} in {s.name!r}. Tables: {[t.name for t in s.tables]}")
    return s, ts[0]


def snapshot(doc):
    return {(s.name, t.name): [[repr(c.value) for c in r] for r in t.rows()] for s in doc.sheets for t in s.tables}


def verify(path, before, password, changed):
    """Reopen the saved file and check that nothing else moved."""
    doc2 = open_numbers(path, password)
    after = snapshot(doc2)
    problems = []
    if set(after) != set(before):
        problems.append("sheet/table list changed")
    for key, rows in before.items():
        rows2 = after.get(key, [])
        for ri, r in enumerate(rows):
            for ci, v in enumerate(r):
                if (key, ri, ci) in changed:
                    continue
                v2 = rows2[ri][ci] if ri < len(rows2) and ci < len(rows2[ri]) else "None"
                if v != v2:
                    problems.append(f"{key[0]}/{key[1]} r{ri + 1}c{ci + 1}: {v} -> {v2}")
    return problems


def default_out(path, tag):
    p = Path(path)
    return str(p.with_name(f"{p.stem}_{tag}{p.suffix}"))


def cmd_set(a):
    from numbers_parser import xl_cell_to_rowcol
    if Path(a.file).suffix.lower() != ".numbers":
        die("editing is available for Numbers files in this version (Pages/Keynote: read and convert only)")
    doc = open_numbers(a.file, a.password)
    before = snapshot(doc)
    s, t = pick_table(doc, a.sheet, a.table)
    changed = set()
    for pair in a.pairs:
        if "=" not in pair:
            die(f"use CELL=value, e.g. B3=1200 (got {pair!r})")
        ref, val = pair.split("=", 1)
        try:
            r, c = xl_cell_to_rowcol(ref.strip().upper())
        except Exception:  # noqa: BLE001
            die(f"bad cell address {ref!r}")
        old = t.cell(r, c).value if r < t.num_rows and c < t.num_cols else None
        t.write(r, c, parse_value(val))
        changed.add(((s.name, t.name), r, c))
        print(f"{s.name}/{t.name} {ref.upper()}: {old!r} -> {val!r}")
    out = a.out or default_out(a.file, "edited")
    doc.save(out, password=a.password) if a.password else doc.save(out)
    probs = verify(out, before, a.password, changed)
    if probs:
        print("WARNING — other cells changed after saving:\n  " + "\n  ".join(probs[:20]))
    else:
        print(f"saved: {out}  (checked: every other cell is unchanged)")


# ── Pages / Keynote text replace ──
def _u16(s):
    return sum(2 if ord(c) > 0xFFFF else 1 for c in s)


def _char_indexes(o, found):
    """Collect every dict that carries a characterIndex inside one text storage."""
    if isinstance(o, dict):
        if isinstance(o.get("characterIndex"), int):
            found.append(o)
        for v in o.values():
            _char_indexes(v, found)
    elif isinstance(o, list):
        for v in o:
            _char_indexes(v, found)
    return found


def _replace_in_storage(st, pairs):
    """Replace text inside one TSWP.StorageArchive and move every style/paragraph marker after it.
    Returns (number replaced, skipped notes)."""
    text = st.get("text")
    if not text or len(text) != 1:
        return 0, []
    t = text[0]
    marks = _char_indexes({k: v for k, v in st.items() if k != "text"}, [])
    hits, skipped = [], []
    for f, rp in pairs:
        start = 0
        while True:
            i = t.find(f, start)
            if i < 0:
                break
            hits.append((i, f, rp))
            start = i + len(f)
    hits.sort()
    # drop overlapping matches (first pair wins)
    clean, end = [], -1
    for h in hits:
        if h[0] >= end:
            clean.append(h)
            end = h[0] + len(h[1])
    done = 0
    for i, f, rp in sorted(clean, reverse=True):
        p, lo, ln = _u16(t[:i]), _u16(f), _u16(rp)
        if any(p < m["characterIndex"] < p + lo for m in marks):
            skipped.append(f"{f!r}: formatting changes inside the word, left as is")
            continue
        for m in marks:
            if m["characterIndex"] >= p + lo:
                m["characterIndex"] += ln - lo
        t = t[:i] + rp + t[i + len(f):]
        done += 1
    text[0] = t
    n16 = _u16(t)
    if any(m["characterIndex"] > n16 for m in marks):
        raise RuntimeError("style markers point past the end of the text")
    return done, skipped


def _context(t, i, n, width=30):
    a, b = max(0, i - width), min(len(t), i + n + width)
    return ("…" if a else "") + clean_text(t[a:b]).replace("\n", " ") + ("…" if b < len(t) else "")


def iwa_replace(path, pairs, out, dry=False):
    from keynote_parser.codec import IWAFile
    import copy
    for f, rp in pairs:
        if not f:
            die("the text to find is empty")
        bad = set("\n\u2028\u2029\ufffc\x04\x05\x0c")
        if bad & set(f) or bad & set(rp):
            die("line breaks can't be added or removed by replace — change one line at a time")
    pkg = Package(path)
    if pkg.format != "iwa":
        die("editing works on Pages/Keynote files saved in 2013 or later (this one is older or locked). "
            "Convert it to Word/PowerPoint instead.")
    allowed = None
    if pkg.kind == "keynote":   # only real slides (and their tables/notes), not the master slide designs
        allowed = {s["file"] for s in keynote_slides(IWA(pkg))}
    changed, report, skipped = {}, [], []
    for n in pkg.list():
        if not n.endswith(".iwa"):
            continue
        if allowed is not None and n not in allowed and not n.startswith("Index/Tables/"):
            continue
        d = IWAFile.from_buffer(pkg.read_any(n), n).to_dict()
        before = copy.deepcopy(d)
        count = 0
        for ch in d.get("chunks", []):
            for a in ch.get("archives", []):
                for o in a.get("objects", []):
                    t = o.get("_pbtype")
                    if t == "TSWP.StorageArchive":
                        old = "".join(o.get("text", []))
                        k, sk = _replace_in_storage(o, pairs)
                        skipped += sk
                        if k:
                            count += k
                            where = "speaker notes" if o.get("kind") == "NOTE" else "text"
                            for f, rp in pairs:
                                j = old.find(f)
                                while j >= 0:
                                    report.append(f"{where}: {_context(old, j, len(f))}  ->  {rp!r}")
                                    j = old.find(f, j + len(f))
                    elif t == "TST.TableDataList" and o.get("listType") == "STRING":
                        for e in o.get("entries", []):
                            s0 = e.get("string")
                            if isinstance(s0, str):
                                s1 = s0
                                for f, rp in pairs:
                                    s1 = s1.replace(f, rp)
                                if s1 != s0:
                                    e["string"] = s1
                                    count += 1
                                    report.append(f"{n} (table cell): {s0[:60]!r} -> {s1[:60]!r}")
        if count:
            changed[n] = (d, before)
    print(f"{len(report)} place(s) match")
    for r in report[:60]:
        print("  " + r)
    for s_ in sorted(set(skipped))[:10]:
        print("  skipped " + s_)
    if dry or not changed:
        return
    new_bytes = {n: IWAFile.from_dict(d).to_buffer() for n, (d, _) in changed.items()}
    write_package(pkg, out, new_bytes)
    # verify: reopen, every changed file decodes, and the untouched files are byte-identical
    pkg2 = Package(out)
    for n in pkg.list():
        if n in new_bytes:
            IWAFile.from_buffer(pkg2.read_any(n), n)
        elif pkg2.read_any(n) != pkg.read_any(n):
            print(f"WARNING: {n} changed unexpectedly")
    r = load(out)
    if r.get("damaged"):
        print("WARNING: saved file has unreadable parts: " + ", ".join(r["damaged"]))
    else:
        print(f"saved: {out}  (checked: reopens cleanly; {len(pkg.list()) - len(new_bytes)} other parts are byte-for-byte unchanged)")
        print("note: the small preview picture inside the file still shows the old text until it is saved again on a Mac")


def write_package(pkg, out, new_bytes):
    """Write a copy of the package with some inner files replaced. Always writes a single zip file."""
    out = Path(out)
    tmp = out.with_suffix(out.suffix + ".tmp")
    inner_new = {}
    if pkg._inner is not None:
        inner_new = {n: b for n, b in new_bytes.items() if n in pkg._inner.namelist()}
    with zipfile.ZipFile(tmp, "w") as z:
        if pkg._zip is not None:
            for info in pkg._zip.infolist():
                name = info.filename
                rel = name[len(pkg.prefix):] if name.startswith(pkg.prefix) else name
                if info.is_dir():
                    continue
                if rel == "Index.zip" and inner_new:
                    buf = io.BytesIO()
                    with zipfile.ZipFile(buf, "w") as iz:
                        for ii in pkg._inner.infolist():
                            zi2 = zipfile.ZipInfo(ii.filename, date_time=ii.date_time)
                            zi2.compress_type = ii.compress_type
                            iz.writestr(zi2, inner_new.get(ii.filename, pkg._inner.read(ii.filename)))
                    data = buf.getvalue()
                else:
                    data = new_bytes.get(rel) if rel in new_bytes and rel not in inner_new else pkg._zip.read(name)
                zi = zipfile.ZipInfo(rel, date_time=info.date_time)
                zi.compress_type = info.compress_type
                zi.external_attr = info.external_attr
                z.writestr(zi, data)
        else:
            for rel in pkg.names:
                r2 = rel[len(pkg.prefix):] if rel.startswith(pkg.prefix) else rel
                data = new_bytes.get(r2) if r2 in new_bytes else (pkg.path / rel).read_bytes()
                z.writestr(zipfile.ZipInfo(r2), data, compress_type=zipfile.ZIP_STORED)
    tmp.replace(out)


def cmd_replace(a):
    if len(a.pairs) % 2:
        die("give pairs: FIND REPLACE [FIND2 REPLACE2 ...]")
    pairs = list(zip(a.pairs[::2], a.pairs[1::2]))
    if Path(a.file).suffix.lower() in (".pages", ".key"):
        iwa_replace(a.file, pairs, a.out or default_out(a.file, "edited"), a.dry)
        return
    doc = open_numbers(a.file, a.password)
    before = snapshot(doc)
    changed, hits = set(), []
    for s in doc.sheets:
        for t in s.tables:
            for ri, row in enumerate(t.rows()):
                for ci, c in enumerate(row):
                    if not isinstance(c.value, str) or c.formula:
                        continue
                    if type(c).__name__ == "RichTextCell":
                        continue
                    new = c.value
                    for f, rp in pairs:
                        new = new.replace(f, rp)
                    if new != c.value:
                        hits.append(f"{s.name}/{t.name} r{ri + 1}c{ci + 1}: {c.value!r} -> {new!r}")
                        if not a.dry:
                            t.write(ri, ci, new)
                        changed.add(((s.name, t.name), ri, ci))
    print(f"{len(hits)} cell(s) match")
    for h in hits[:50]:
        print("  " + h)
    if a.dry or not hits:
        return
    out = a.out or default_out(a.file, "edited")
    doc.save(out, password=a.password) if a.password else doc.save(out)
    probs = verify(out, before, a.password, changed)
    print("WARNING:\n  " + "\n  ".join(probs[:20]) if probs else f"saved: {out}  (checked: every other cell is unchanged)")


def cmd_new(a):
    need("numbers_parser", "numbers-parser")
    from numbers_parser import Document
    _patch_numbers_parser()
    src = Path(a.data)
    sheets = []
    if src.suffix.lower() in (".xlsx", ".xlsm"):
        need("openpyxl", "openpyxl")
        import openpyxl
        wb = openpyxl.load_workbook(src, data_only=True)
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            while rows and all(v is None for v in rows[-1]):
                rows.pop()
            sheets.append((ws.title, rows or [[None]]))
        had_formula = False
        wb2 = openpyxl.load_workbook(src)
        for ws in wb2.worksheets:
            for r in ws.iter_rows():
                for c in r:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        had_formula = True
    elif src.suffix.lower() in (".csv", ".tsv"):
        raw = src.read_bytes()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        rows = list(csv.reader(io.StringIO(text), delimiter="\t" if src.suffix.lower() == ".tsv" else ","))
        rows = [[parse_value(v) if v != "" else None for v in r] for r in rows]
        sheets.append((src.stem[:30] or "Sheet 1", rows or [[None]]))
        had_formula = False
    else:
        die("give an .xlsx or .csv file")
    doc = Document(num_header_rows=1, num_header_cols=0,
                   num_rows=max(len(sheets[0][1]), 1), num_cols=max(max(len(r) for r in sheets[0][1]), 1))
    for i, (name, rows) in enumerate(sheets):
        nr, nc = max(len(rows), 1), max(max((len(r) for r in rows), default=1), 1)
        if i == 0:
            doc.sheets[0].name = name
            t = doc.sheets[0].tables[0]
        else:
            doc.add_sheet(name, "Table 1", num_rows=nr, num_cols=nc)
            t = doc.sheets[-1].tables[0]
        for ri, r in enumerate(rows):
            for ci, v in enumerate(r):
                if v is None:
                    continue
                if isinstance(v, dt.date) and not isinstance(v, dt.datetime):
                    v = dt.datetime(v.year, v.month, v.day)
                if isinstance(v, dt.time):
                    v = v.strftime("%H:%M:%S")
                t.write(ri, ci, v)
    out = a.out or str(src.with_suffix(".numbers"))
    doc.save(out)
    chk = open_numbers(out)
    print(f"saved: {out}  ({len(chk.sheets)} sheet(s), first table {chk.sheets[0].tables[0].num_rows} rows)")
    if had_formula:
        print("note: Excel formulas were saved as their last calculated values (Numbers formulas cannot be "
              "written yet; a formula with no saved result is left empty)")


def main():
    ap = argparse.ArgumentParser(description="Apple iWork file tool (Pages, Numbers, Keynote)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("info"); s.add_argument("file"); s.add_argument("--password"); s.set_defaults(fn=cmd_info)
    s = sub.add_parser("read"); s.add_argument("file"); s.add_argument("--out"); s.add_argument("--password")
    s.add_argument("--formulas", action="store_true", help="show Numbers formulas next to values"); s.set_defaults(fn=cmd_read)
    s = sub.add_parser("preview"); s.add_argument("file"); s.add_argument("--out"); s.set_defaults(fn=cmd_preview)
    s = sub.add_parser("images"); s.add_argument("file"); s.add_argument("--out")
    s.add_argument("--all", action="store_true", help="also theme/template pictures"); s.set_defaults(fn=cmd_images)
    s = sub.add_parser("convert"); s.add_argument("file"); s.add_argument("out"); s.add_argument("--password")
    s.add_argument("--values-only", action="store_true", help="Excel: write values, not formulas")
    s.set_defaults(fn=cmd_convert, formulas=False)
    s = sub.add_parser("set"); s.add_argument("file"); s.add_argument("pairs", nargs="+"); s.add_argument("--sheet")
    s.add_argument("--table"); s.add_argument("--out"); s.add_argument("--password"); s.set_defaults(fn=cmd_set)
    s = sub.add_parser("replace"); s.add_argument("file"); s.add_argument("pairs", nargs="+"); s.add_argument("--out")
    s.add_argument("--password"); s.add_argument("--dry", action="store_true"); s.set_defaults(fn=cmd_replace)
    s = sub.add_parser("new"); s.add_argument("data"); s.add_argument("--out"); s.set_defaults(fn=cmd_new)
    a = ap.parse_args()
    if not hasattr(a, "password"):
        a.password = None
    if not hasattr(a, "formulas"):
        a.formulas = False
    a.fn(a)


if __name__ == "__main__":
    main()
